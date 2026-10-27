#!/usr/bin/env python3
"""Validate a BRD JSON file and render it to PDF and Word in the sample BRD layout.

Usage:
  brd_render.py check  <brd.json>          # validate only; exit 1 on problems
  brd_render.py render <brd.json> [outdir] # validate, then write .html/.pdf/.docx
  brd_render.py coverage <brd.json> <source.pdf|summary.md>  # how much source text made it into the BRD
  brd_render.py settled <run folder>       # OIs the prototype fully answered (leave out of the BRD) and partly answered

Run with the project venv (needs python-docx):
  .claude/skills/req-team/.venv/bin/python .claude/skills/brd-format/brd_render.py render <brd.json>

The layout follows input/BRD-Formatter/RD_Team Store M-1.0(Landing Page).docx.pdf:
header table, Versions table, Functional Requirements rows (ID | name | details,
with child rows that pair a screenshot with bullets), Phased-in Requirements and
References.
"""
import base64
import html
import json
import mimetypes
import os
import re
import subprocess
import sys
import tempfile

CHROME = "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome"
HEADER_FIELDS = ["project_name", "type", "stakeholder", "us_uiux", "ba_author",
                 "project_manager", "development_team", "sqa_team", "other_teams",
                 "priority", "brand"]


# ---------------------------------------------------------------- validation

def _items_problems(items, where):
    problems = []
    if not isinstance(items, list) or not items:
        return [f"{where}: needs at least one item."]
    for i, it in enumerate(items, 1):
        if isinstance(it, str):
            it = {"text": it}
        if not isinstance(it, dict):
            problems.append(f"{where} item {i}: must be a string or an object.")
            continue
        if "table" in it:
            # {"table": {"columns": [...], "rows": [[cell, ...], ...]}}: a source table kept as a table
            tb = it["table"] or {}
            cols, rows = tb.get("columns") or [], tb.get("rows") or []
            if not cols or not rows:
                problems.append(f"{where} item {i}: table needs 'columns' and 'rows'.")
            for r_no, r in enumerate(rows, 1):
                if not isinstance(r, list) or len(r) != len(cols):
                    problems.append(f"{where} item {i}: table row {r_no} needs {len(cols)} cells.")
            continue
        if not (it.get("text") or it.get("label")):
            problems.append(f"{where} item {i}: empty (needs 'text' or 'label').")
        if "items" in it:
            problems += _items_problems(it["items"], f"{where} item {i}")
    return problems


def check(brd, base_dir="."):
    """Return a list of problems with the BRD structure."""
    p = []
    mod = brd.get("module") or {}
    mid = mod.get("id", "")
    m = re.fullmatch(r"M-(\d+)\.(\d+)", mid)
    if not m:
        return [f"module.id must look like 'M-2.0' or 'M-4.1' (got '{mid}')."]
    # M-2.0 numbers its rows FR-2.k; a sub-module such as M-4.1 numbers them FR-4.1.k
    n = m.group(1) if m.group(2) == "0" else f"{m.group(1)}.{m.group(2)}"
    if not mod.get("name"):
        p.append("module.name is missing.")
    if mod.get("items") is not None:
        p += _items_problems(mod["items"], mid)
    hdr = brd.get("header") or {}
    if not hdr.get("project_name"):
        p.append("header.project_name is missing.")
    for k in hdr:
        if k not in HEADER_FIELDS:
            p.append(f"header.{k} is not a known field ({', '.join(HEADER_FIELDS)}).")
    if not brd.get("versions"):
        p.append("versions needs at least one row.")
    ids, reqs = [], []
    if brd.get("sections"):
        for s_no, sec in enumerate(brd["sections"], 1):
            if sec.get("id") != f"FR-{s_no}":
                p.append(f"section {s_no}: id should be 'FR-{s_no}' (got '{sec.get('id')}').")
            if not sec.get("name"):
                p.append(f"FR-{s_no}: section name is missing.")
            if sec.get("items") is not None:
                p += _items_problems(sec["items"], f"FR-{s_no}")
            if not sec.get("requirements"):
                p.append(f"FR-{s_no}: section has no requirements.")
            ids.append(sec.get("id", ""))
            reqs += [(f"FR-{s_no}", k, r) for k, r in enumerate(sec.get("requirements") or [], 1)]
    else:
        reqs = [(f"FR-{n}", k, r) for k, r in enumerate(brd.get("requirements") or [], 1)]
    if not reqs:
        p.append("requirements is empty.")
    for prefix, k, r in reqs:
        rid = r.get("id", "")
        if rid != f"{prefix}.{k}":
            p.append(f"requirement {k} of {prefix}: id should be '{prefix}.{k}' (got '{rid}'). IDs run in order with no gaps.")
        ids.append(rid)
        if not r.get("name"):
            p.append(f"{rid}: name is missing.")
        blocks = r.get("blocks") or []
        if not blocks:
            p.append(f"{rid}: needs at least one block.")
        for b in blocks:
            p += _items_problems(b.get("items"), f"{rid} block '{b.get('heading', '')}'")
        for c, ch in enumerate(r.get("children") or [], 1):
            cid = ch.get("id", "")
            if cid != f"{rid}.{c}":
                p.append(f"{rid} child {c}: id should be '{rid}.{c}' (got '{cid}').")
            p += _items_problems(ch.get("items"), cid or f"{rid} child {c}")
            shot = ch.get("screenshot")
            if shot and not os.path.exists(os.path.join(base_dir, shot)) and not os.path.exists(shot):
                p.append(f"{cid}: screenshot file not found: {shot}")
    idset = set(ids)
    for _, _, r in reqs:
        nxt = r.get("next")
        if nxt and nxt not in idset:
            p.append(f"{r.get('id')}: next points to unknown '{nxt}'.")
    for ph in brd.get("phased") or []:
        for ref in re.findall(r"FR-[\d.]+", ph.get("fr", "")):
            if ref.rstrip(".") not in idset and not any(ref.startswith(i + ".") for i in idset):
                p.append(f"phased row refers to unknown {ref}.")
    if not brd.get("references"):
        p.append("references is empty.")
    return p


# ---------------------------------------------------------------- shared helpers

def sep(label):
    """Colon after a lead-in label, unless it already ends in punctuation ("Render conditions.", "Card 1 —")."""
    return "" if label.rstrip()[-1:] in ".:;?!—–-" else ":"


def groups(brd):
    """[(section or None, requirements)]: a consolidated BRD has sections FR-1, FR-2 …"""
    if brd.get("sections"):
        return [(sec, sec.get("requirements") or []) for sec in brd["sections"]]
    return [(None, brd.get("requirements") or [])]


def norm(items):
    return [{"text": it} if isinstance(it, str) else it for it in items or []]


def img_data_uri(path):
    mime = mimetypes.guess_type(path)[0] or "image/png"
    with open(path, "rb") as f:
        return f"data:{mime};base64,{base64.b64encode(f.read()).decode()}"


def resolve(path, base_dir):
    if not path:
        return None
    for cand in (path, os.path.join(base_dir, path)):
        if os.path.exists(cand):
            return cand
    return None


# ---------------------------------------------------------------- HTML / PDF

CSS = """
@page { size: A4; margin: 14mm 12mm; }
* { box-sizing: border-box; }
body { font: 9.5pt/1.38 Calibri, Carlito, "Segoe UI", Arial, sans-serif; color: #000; margin: 0; }
h1 { text-align: center; font-size: 11pt; font-weight: 700; margin: 0 0 8px; }
table { width: 100%; border-collapse: collapse; margin: 0 0 10px; table-layout: fixed; }
td, th { border: 1px solid #c8c8c8; padding: 5px 7px; vertical-align: top; text-align: left; }
.lbl { background: #d9d9d9; font-weight: 700; font-size: 8.8pt; }
.sub { background: #ededed; font-weight: 700; font-size: 8.8pt; text-align: center; }
.blank td { height: 22px; }
.sec { background: #d9d9d9; font-weight: 700; font-size: 8.8pt; }
.fr td { page-break-inside: auto; }
.secrow td { background: #f2f2f2; }
.fr .id { width: 11%; font-weight: 700; font-size: 8.8pt; }
.fr .nm { width: 26%; }
.fr .nm b { display: block; }
.fr .el { color: #666; font-size: 8pt; display: block; }
.fr .ep { font-style: italic; font-weight: 700; display: block; }
.fr .lnk { color: #1155cc; text-decoration: underline; font-size: 8.5pt; }
ul { margin: 0 0 4px; padding-left: 15px; }
li { margin: 1px 0; }
.bh { font-weight: 700; margin: 2px 0; }
.next { color: #555; font-size: 8.5pt; margin-top: 3px; }
.child { padding: 8px 10px 10px; }
.cid { font-weight: 700; font-size: 8.5pt; color: #555; margin-bottom: 4px; }
.shot { border: 1px dashed #9aa; background: #f6f8f9; color: #678; text-align: center;
        padding: 26px 10px; font-size: 8.5pt; margin-bottom: 6px; }
.shot img { max-width: 100%; }
.sl { font-weight: 700; margin: 2px 0; }
li.tli { list-style: none; margin: 4px 0 6px -15px; }
table.tt { table-layout: auto; margin: 0; font-size: 8.5pt; }
table.tt th { background: #ededed; font-weight: 700; }
table.tt th, table.tt td { padding: 3px 5px; }
table.tt.idcol td:first-child { white-space: nowrap; }
table.tt:not(.idcol) td:first-child { width: 22%; }
table.tt.wide { font-size: 7.8pt; }
table.tt thead { display: table-header-group; }
tr { page-break-inside: avoid; }
tr.big { page-break-inside: auto; }
"""


def h_table(tb):
    e = html.escape
    head = "".join(f"<th>{e(c)}</th>" for c in tb["columns"])
    body = "".join("<tr>" + "".join(f"<td>{e(str(c))}</td>" for c in r) + "</tr>" for r in tb["rows"])
    wide = " wide" if len(tb["columns"]) > 4 else ""
    # keep short ID columns (C-01, BR-7) on one line; let longer first columns (glossary terms) wrap
    idcol = " idcol" if all(len(str(r[0])) <= 12 for r in tb["rows"] if r) else ""
    # thead repeats the header row when a long table runs onto the next page
    return f"<table class='tt{wide}{idcol}'><thead><tr>{head}</tr></thead><tbody>{body}</tbody></table>"


def h_items(items):
    out = []
    for it in norm(items):
        if it.get("table"):
            out.append(f"<li class=tli>{h_table(it['table'])}</li>")
            continue
        label, text = it.get("label"), it.get("text", "")
        body = (f"<b>{html.escape(label)}{sep(label)}</b> " if label and text else
                f"<b>{html.escape(label)}</b>" if label else "") + html.escape(text)
        sub = h_items(it["items"]) if it.get("items") else ""
        out.append(f"<li>{body}{sub}</li>")
    return "<ul>" + "".join(out) + "</ul>"


def h_blocks(blocks):
    out = []
    for b in blocks:
        if b.get("heading"):
            out.append(f"<div class='bh'>{html.escape(b['heading'])}</div>")
        out.append(h_items(b.get("items")))
    return "".join(out)


def to_html(brd, base_dir):
    e = html.escape
    mod, hdr = brd["module"], brd["header"]
    g = lambda k: e(hdr.get(k, "") or "")
    parts = [f"<h1>REQUIREMENTS DOCUMENT — {e(mod['id'])}</h1>"]
    parts.append(f"""<table><colgroup><col style="width:34%"><col style="width:34%"><col style="width:32%"></colgroup>
<tr><td class=lbl colspan=2>PROJECT NAME:</td><td class=lbl>Type</td></tr>
<tr><td colspan=2><b>{g('project_name')}</b></td><td>{g('type')}</td></tr>
<tr><td class=lbl>STAKEHOLDER</td><td class=lbl colspan=2>US UI/UX:</td></tr>
<tr><td>{g('stakeholder')}</td><td colspan=2>{g('us_uiux')}</td></tr>
<tr><td class=lbl colspan=3>BA AUTHOR:</td></tr><tr><td colspan=3>{g('ba_author')}</td></tr>
<tr><td class=lbl colspan=3>PROJECT MANAGER:</td></tr><tr><td colspan=3>{g('project_manager')}</td></tr>
<tr><td class=lbl colspan=3>DEVELOPMENT TEAM: (INCLUDING UI/UX)</td></tr><tr><td colspan=3>{g('development_team') or '&nbsp;'}</td></tr>
<tr><td class=lbl colspan=3>SQA TEAM:</td></tr><tr><td colspan=3>{g('sqa_team') or '&nbsp;'}</td></tr>
<tr><td class=lbl colspan=3>OTHER TEAMS:</td></tr><tr><td colspan=3>{g('other_teams') or '&nbsp;'}</td></tr>
<tr><td class=lbl>PRIORITY:</td><td class=lbl colspan=2>BRAND:</td></tr>
<tr><td>{g('priority') or '&nbsp;'}</td><td colspan=2>{g('brand') or '&nbsp;'}</td></tr>
</table>""")
    vers = brd["versions"]
    rows = "".join(f"<tr><td>{e(v.get('version',''))}</td><td>{e(v.get('br_fr',''))}</td>"
                   f"<td>{e(v.get('remarks',''))}</td><td>{e(v.get('date',''))}</td></tr>" for v in vers)
    rows += "<tr class=blank><td></td><td></td><td></td><td></td></tr>" * max(0, 3 - len(vers))
    parts.append(f"""<table><tr><td class=lbl colspan=4>VERSIONS<br><span style="font-weight:400">{e(vers[0].get('version',''))} -</span></td></tr>
<tr><td class=sub>VERSION NO.</td><td class=sub>BR/FR NO.</td><td class=sub>REVISION REMARKS</td><td class=sub>LAST UPDATED DATE</td></tr>{rows}</table>""")

    fr = ["<table class=fr><colgroup><col style='width:11%'><col style='width:26%'><col style='width:63%'></colgroup>"
          "<tr><td class=sec colspan=3>FUNCTIONAL REQUIREMENTS (FR)</td></tr>"]
    fr.append(f"<tr><td class=id>{e(mod['id'])}</td><td class=nm><b>{e(mod['name'])}</b>"
              f"{'<span class=el>' + e(mod['element']) + '</span>' if mod.get('element') else ''}</td>"
              f"<td>{h_items(mod.get('items')) if mod.get('items') else '—'}</td></tr>")
    for sec, sreqs in groups(brd):
      if sec:
        fr.append(f"<tr class='secrow big'><td class=id>{e(sec['id'])}</td><td class=nm><b>{e(sec['name'])}</b>"
                  f"{'<span class=el>' + e(sec['element']) + '</span>' if sec.get('element') else ''}</td>"
                  f"<td>{h_items(sec.get('items')) if sec.get('items') else '—'}</td></tr>")
      for r in sreqs:
        nm = f"<b>{e(r['name'])}</b>"
        if r.get("subtitle"):
            nm += f"<span class=ep>{e(r['subtitle'])}</span>"
        if r.get("element"):
            nm += f"<span class=el>{e(r['element'])}</span>"
        if r.get("link"):
            nm += f"<span class=lnk>{e(r['link'])}</span>"
        det = h_blocks(r["blocks"])
        if r.get("next"):
            det += f"<div class=next>Next: {e(r['next'])}</div>"
        fr.append(f"<tr class=big><td class=id>{e(r['id'])}</td><td class=nm>{nm}</td><td>{det}</td></tr>")
        for ch in r.get("children") or []:
            shot = resolve(ch.get("screenshot"), base_dir)
            cap = e(ch.get("caption", ""))
            box = (f"<div class=shot style='padding:4px'><img src='{img_data_uri(shot)}'></div>" if shot else
                   f"<div class=shot>Screenshot placeholder — {cap or 'attach final design'}</div>")
            fr.append(f"<tr><td></td><td colspan=2 class=child><div class=cid>{e(ch['id'])}</div>{box}"
                      f"<div class=sl>//</div>{h_items(ch['items'])}</td></tr>")
    fr.append("</table>")
    parts.append("".join(fr))

    ph = brd.get("phased") or []
    prow = "".join(f"<tr><td>{e(x.get('fr',''))}</td><td>{e(x.get('requirement',''))}</td>"
                   f"<td>{e(x.get('phase',''))}</td><td>{e(x.get('remarks',''))}</td></tr>" for x in ph)
    prow += "<tr class=blank><td></td><td></td><td></td><td></td></tr>" * max(0, 2 - len(ph))
    parts.append(f"""<table><tr><td class=lbl colspan=4>PHASED IN REQUIREMENTS</td></tr>
<tr><td class=sub>FR NO.</td><td class=sub>REQUIREMENT</td><td class=sub>TARGET PHASE</td><td class=sub>REMARKS</td></tr>{prow}</table>""")
    rrow = "".join(f"<tr><td style='width:34%'>{e(x.get('name',''))}</td><td>{e(x.get('location',''))}</td></tr>"
                   for x in brd["references"])
    parts.append(f"<table><colgroup><col style='width:34%'><col style='width:66%'></colgroup>"
                 f"<tr><td class=lbl colspan=2>REFERENCES</td></tr>"
                 f"<tr><td class=lbl style='width:34%'>NAME</td><td class=lbl>LOCATION</td></tr>{rrow}</table>")
    title = e(f"RD_{brd['header'].get('project_name', '')} {mod['id']}")
    return (f"<!doctype html><html><head><meta charset='utf-8'><title>{title}</title>"
            f"<style>{CSS}</style></head><body>{''.join(parts)}</body></html>")


def to_pdf(html_text, pdf_path):
    with tempfile.NamedTemporaryFile("w", suffix=".html", delete=False, encoding="utf-8") as tmp:
        tmp.write(html_text)
    try:
        subprocess.run([CHROME, "--headless", "--disable-gpu", "--no-pdf-header-footer",
                        f"--print-to-pdf={os.path.abspath(pdf_path)}", f"file://{tmp.name}"],
                       check=True, capture_output=True)
    finally:
        os.unlink(tmp.name)


# ---------------------------------------------------------------- DOCX

def to_docx(brd, docx_path, base_dir):
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(1.4)
    sec.top_margin = sec.bottom_margin = Cm(1.4)
    st = doc.styles["Normal"]
    st.font.name, st.font.size = "Calibri", Pt(9.5)
    st.paragraph_format.space_after = Pt(0)
    width = Cm(18.2)

    def shade(cell, color):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), color)
        tcPr.append(shd)

    def borders(table):
        tblPr = table._tbl.tblPr
        b = OxmlElement("w:tblBorders")
        for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
            el = OxmlElement(f"w:{side}")
            el.set(qn("w:val"), "single"); el.set(qn("w:sz"), "4"); el.set(qn("w:color"), "C8C8C8")
            b.append(el)
        tblPr.append(b)

    def table(cols, widths):
        t = doc.add_table(rows=0, cols=cols)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        borders(t)
        t._widths = widths
        return t

    def row(t, texts, fill=None, bold=False, center=False, merge=None):
        cells = t.add_row().cells
        for i, w in enumerate(t._widths):
            cells[i].width = w
        if merge:
            a, z = merge
            cells[a].merge(cells[z])
        for i, txt in enumerate(texts):
            if txt is None:
                continue
            c = cells[i]
            c.paragraphs[0].text = ""
            r = c.paragraphs[0].add_run(txt)
            r.bold = bold
            if center:
                c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            if fill:
                shade(c, fill)
        if fill and merge:
            shade(cells[merge[0]], fill)
        return cells

    def cell_table(cell, tb):
        t = cell.add_table(rows=0, cols=len(tb["columns"]))
        borders(t)
        for r_no, texts in enumerate([tb["columns"]] + tb["rows"]):
            cells = t.add_row().cells
            for c, txt in zip(cells, texts):
                run = c.paragraphs[0].add_run(str(txt))
                run.font.size = Pt(8.5)
                if r_no == 0:
                    run.bold = True
                    shade(c, "EDEDED")
        cell.add_paragraph()  # Word needs a paragraph after a nested table

    def bullets(cell, items, level=0):
        for it in norm(items):
            if it.get("table"):
                cell_table(cell, it["table"])
                continue
            p = cell.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
            if it.get("label"):
                p.add_run(it["label"] + (sep(it["label"]) + " " if it.get("text") else "")).bold = True
            if it.get("text"):
                p.add_run(it["text"])
            if it.get("items"):
                bullets(cell, it["items"], level + 1)

    def clear_first(cell):
        if cell.paragraphs and not cell.paragraphs[0].text:
            cell._tc.remove(cell.paragraphs[0]._p)

    mod, hdr = brd["module"], brd["header"]
    h = doc.add_paragraph()
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    hr = h.add_run(f"REQUIREMENTS DOCUMENT — {mod['id']}")
    hr.bold, hr.font.size = True, Pt(11)

    g = lambda k: hdr.get(k, "") or ""
    w3 = [Cm(6.2), Cm(6.2), Cm(5.8)]
    t = table(3, w3)
    row(t, ["PROJECT NAME:", None, "Type"], fill="D9D9D9", bold=True, merge=(0, 1))
    c = row(t, [g("project_name"), None, g("type")], merge=(0, 1)); c[0].paragraphs[0].runs[0].bold = True
    row(t, ["STAKEHOLDER", "US UI/UX:", None], fill="D9D9D9", bold=True, merge=(1, 2))
    row(t, [g("stakeholder"), g("us_uiux"), None], merge=(1, 2))
    for lbl, key in (("BA AUTHOR:", "ba_author"), ("PROJECT MANAGER:", "project_manager"),
                     ("DEVELOPMENT TEAM: (INCLUDING UI/UX)", "development_team"),
                     ("SQA TEAM:", "sqa_team"), ("OTHER TEAMS:", "other_teams")):
        row(t, [lbl, None, None], fill="D9D9D9", bold=True, merge=(0, 2))
        row(t, [g(key), None, None], merge=(0, 2))
    row(t, ["PRIORITY:", "BRAND:", None], fill="D9D9D9", bold=True, merge=(1, 2))
    row(t, [g("priority"), g("brand"), None], merge=(1, 2))
    doc.add_paragraph()

    w4 = [Cm(4.55)] * 4
    vers = brd["versions"]
    t = table(4, w4)
    row(t, [f"VERSIONS\n{vers[0].get('version', '')} -", None, None, None], fill="D9D9D9", bold=True, merge=(0, 3))
    row(t, ["VERSION NO.", "BR/FR NO.", "REVISION REMARKS", "LAST UPDATED DATE"], fill="EDEDED", bold=True, center=True)
    for v in vers:
        row(t, [v.get("version", ""), v.get("br_fr", ""), v.get("remarks", ""), v.get("date", "")])
    for _ in range(max(0, 3 - len(vers))):
        row(t, ["", "", "", ""])
    doc.add_paragraph()

    wf = [Cm(2.0), Cm(4.7), Cm(11.5)]
    t = table(3, wf)
    row(t, ["FUNCTIONAL REQUIREMENTS (FR)", None, None], fill="D9D9D9", bold=True, merge=(0, 2))
    c = row(t, [mod["id"], mod["name"], ""])
    c[0].paragraphs[0].runs[0].bold = True; c[1].paragraphs[0].runs[0].bold = True
    if mod.get("element"):
        rr = c[1].add_paragraph().add_run(mod["element"]); rr.font.size = Pt(8); rr.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    if mod.get("items"):
        bullets(c[2], mod["items"]); clear_first(c[2])
    else:
        c[2].paragraphs[0].text = "—"
    for sec, sreqs in groups(brd):
      if sec:
        c = row(t, [sec["id"], sec["name"], ""], fill="F2F2F2")
        c[0].paragraphs[0].runs[0].bold = True; c[1].paragraphs[0].runs[0].bold = True
        if sec.get("items"):
            bullets(c[2], sec["items"]); clear_first(c[2])
        else:
            c[2].paragraphs[0].text = "—"
      for r in sreqs:
        c = row(t, [r["id"], r["name"], ""])
        c[0].paragraphs[0].runs[0].bold = True; c[1].paragraphs[0].runs[0].bold = True
        if r.get("subtitle"):
            rr = c[1].add_paragraph().add_run(r["subtitle"]); rr.bold = rr.italic = True
        if r.get("element"):
            rr = c[1].add_paragraph().add_run(r["element"]); rr.font.size = Pt(8); rr.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
        if r.get("link"):
            rr = c[1].add_paragraph().add_run(r["link"]); rr.underline = True; rr.font.color.rgb = RGBColor(0x11, 0x55, 0xCC)
        for b in r["blocks"]:
            if b.get("heading"):
                c[2].add_paragraph().add_run(b["heading"]).bold = True
            bullets(c[2], b.get("items"))
        if r.get("next"):
            rr = c[2].add_paragraph().add_run(f"Next: {r['next']}"); rr.font.size = Pt(8.5); rr.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
        clear_first(c[2])
        for ch in r.get("children") or []:
            c = row(t, ["", "", None], merge=(1, 2))
            cell = c[1]
            cell.paragraphs[0].add_run(ch["id"]).bold = True
            shot = resolve(ch.get("screenshot"), base_dir)
            if shot:
                cell.add_paragraph().add_run().add_picture(shot, width=Cm(14.5))
            else:
                p = cell.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                rr = p.add_run(f"[ Screenshot placeholder — {ch.get('caption') or 'attach final design'} ]")
                rr.italic = True; rr.font.color.rgb = RGBColor(0x66, 0x77, 0x88)
            cell.add_paragraph().add_run("//").bold = True
            bullets(cell, ch["items"])
    doc.add_paragraph()

    ph = brd.get("phased") or []
    t = table(4, w4)
    row(t, ["PHASED IN REQUIREMENTS", None, None, None], fill="D9D9D9", bold=True, merge=(0, 3))
    row(t, ["FR NO.", "REQUIREMENT", "TARGET PHASE", "REMARKS"], fill="EDEDED", bold=True, center=True)
    for x in ph:
        row(t, [x.get("fr", ""), x.get("requirement", ""), x.get("phase", ""), x.get("remarks", "")])
    for _ in range(max(0, 2 - len(ph))):
        row(t, ["", "", "", ""])
    doc.add_paragraph()

    t = table(2, [Cm(6.2), Cm(12.0)])
    row(t, ["REFERENCES", None], fill="D9D9D9", bold=True, merge=(0, 1))
    row(t, ["NAME", "LOCATION"], fill="D9D9D9", bold=True)
    for x in brd["references"]:
        row(t, [x.get("name", ""), x.get("location", "")])
    doc.save(docx_path)


# ---------------------------------------------------------------- coverage

def _words(text):
    text = text.lower().replace("’", "'").replace("“", '"').replace("”", '"')
    return re.findall(r"[a-z0-9]+", text)


def _strings(node, labels=True, in_row=False):
    """Yield all text. labels: True keeps every label, False drops them all, "rows" drops only
    column labels (labels of items nested under a labelled table row) and keeps row labels."""
    if isinstance(node, str):
        yield node
    elif isinstance(node, dict):
        keep = labels is True or (labels == "rows" and not in_row)
        for k, v in node.items():
            if k == "label" and not keep:
                continue
            yield from _strings(v, labels, in_row=(k == "items" and bool(node.get("label"))))
    elif isinstance(node, list):
        for v in node:
            yield from _strings(v, labels, in_row)


def coverage(brd, source_pdf, n=6, min_gap=12):
    """Share of the source's 6-word sequences found in the BRD, plus the longest missing passages."""
    if not source_pdf.lower().endswith(".md"):
        import fitz
    # Match against the text both with and without labels: column labels between
    # short table cells ("Draft: Allowed  Open: Blocked") would otherwise hide verbatim cell runs.
    have = set()
    for labels in (True, False, "rows"):
        brd_words = _words(" ".join(_strings(brd, labels)))
        have |= {tuple(brd_words[i:i + n]) for i in range(len(brd_words) - n + 1)}
    src, pages = [], []
    if source_pdf.lower().endswith(".md"):
        # Markdown source (an implementation summary): "pages" are line numbers
        with open(source_pdf, encoding="utf-8") as f:
            chunks = enumerate(f.read().splitlines(), 1)
    else:
        chunks = ((pno, page.get_text()) for pno, page in enumerate(fitz.open(source_pdf), 1))
    for pno, text in chunks:
        w = _words(text)
        src += w
        pages += [pno] * len(w)
    total = max(1, len(src) - n + 1)
    hit = [tuple(src[i:i + n]) in have for i in range(total)]
    gaps, start = [], None
    for i, h in enumerate(hit + [True]):
        if not h and start is None:
            start = i
        elif h and start is not None:
            if i - start >= min_gap:
                gaps.append((pages[start], " ".join(src[start:i + n - 1])))
            start = None
    return sum(hit) / total, gaps


# ---------------------------------------------------------------- main

def settled(run):
    """Split the summary's OIs by prototype status: an OI is settled when every Q-ID it maps to is Answered."""
    def read(name):
        p = os.path.join(run, name)
        return open(p, encoding="utf-8").read() if os.path.exists(p) else ""
    questions, answers = read("clarification-questions.md"), read("prototype-answers.md")
    m = re.search(r"^## Answers\s*$(.*?)(?=^## |\Z)", answers, re.M | re.S)
    status = {q: s.strip() for q, s in re.findall(r"^\|\s*(Q-\d+)\s*\|[^|]*\|([^|]*)\|",
                                                     m.group(1) if m else "", re.M)}
    oi_qs = {}
    for line in questions.splitlines():
        q = re.match(r"\|\s*(Q-\d+)\s*\|", line)
        if q:
            for oi in set(re.findall(r"OI-\d+", line)):
                oi_qs.setdefault(oi, []).append(q.group(1))
    out = {"settled": {}, "partial": {}}
    for oi in sorted(oi_qs, key=lambda s: int(s[3:])):
        qs = oi_qs[oi]
        still_open = [q for q in qs if status.get(q) != "Answered"]
        if not still_open:
            out["settled"][oi] = qs
        elif len(still_open) < len(qs):
            out["partial"][oi] = still_open
    return out


def main():
    if len(sys.argv) == 3 and sys.argv[1] == "settled":
        print(json.dumps(settled(sys.argv[2]), indent=1))
        return
    if len(sys.argv) < 3 or sys.argv[1] not in ("check", "render", "coverage"):
        print(__doc__)
        sys.exit(2)
    path = sys.argv[2]
    base_dir = os.path.dirname(os.path.abspath(path))
    with open(path, encoding="utf-8") as f:
        brd = json.load(f)
    if sys.argv[1] == "coverage":
        ratio, gaps = coverage(brd, sys.argv[3])
        print(f"Coverage: {ratio:.1%} of the source text appears in the BRD "
              f"({'PASS' if ratio >= 0.95 and not gaps else 'REVIEW'}; target 95% and no missing passages).")
        where = "line " if sys.argv[3].lower().endswith(".md") else "p."
        for page, text in gaps:
            print(f"- missing (source {where}{page}): {text[:300]}")
        return
    problems = check(brd, base_dir)
    if problems:
        print("BRD check failed:\n- " + "\n- ".join(problems))
        sys.exit(1)
    if sys.argv[1] == "check":
        print("BRD check passed.")
        return
    outdir = sys.argv[3] if len(sys.argv) > 3 else base_dir
    # "Order Management (Orders Tab)" -> "Order Management — Orders Tab" so the name has one set of brackets
    name = re.sub(r"\s*\((.+)\)\s*$", r" — \1", brd["module"]["name"])
    stem = f"RD_{brd.get('file_prefix') or 'Team Store'} {brd['module']['id']} ({name})"
    stem = re.sub(r'[/\\:*?"<>|]', "-", stem)
    html_text = to_html(brd, base_dir)
    html_path = os.path.join(outdir, stem + ".html")
    with open(html_path, "w", encoding="utf-8") as f:
        f.write(html_text)
    to_pdf(html_text, os.path.join(outdir, stem + ".pdf"))
    to_docx(brd, os.path.join(outdir, stem + ".docx"), base_dir)
    for ext in (".pdf", ".docx", ".html"):
        print(os.path.join(outdir, stem + ext))


if __name__ == "__main__":
    main()
